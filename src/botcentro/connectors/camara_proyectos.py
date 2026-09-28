"""Conector SRC-06: proyectos de ley de la Cámara (docs/fuentes/SRC-06.md).

Fuente: https://www.camara.gov.co/proyectos-de-ley/ (licencia publicada: uso, transformación,
redistribución y extracción permitidos citando la fuente). La página usa dos acciones de
`admin-ajax.php` (permitido por robots.txt) con un nonce público incluido en la misma página:

* `download_proyectos_ley_xlsx`: listado completo en XLSX (una petición): números de Cámara y
  Senado, fechas de radicación, título, objeto, tipo, autores, estado, comisiones y legislatura;
* `get_proyectos_ley_page`: listado paginado JSON; aporta el ID de perfil de cada representante
  autor, necesario para no fusionar homónimos por nombre (SRS-F08).

Contradicciones de la fuente (p. ej. un número con sufijo «C» en la columna Senado) no se
corrigen: el candidato va a cuarentena con su motivo.
"""

from __future__ import annotations

import io
import json
import re
from collections.abc import Mapping, Sequence
from datetime import date, datetime
from typing import Any

from botcentro.connectors.base import (
    Cursor,
    DiscoveredItem,
    DiscoverPage,
    Issue,
    IssueSeverity,
    NormalizedCandidate,
    ParseResult,
    ValidationReport,
)
from botcentro.domain.dates import PartialDate
from botcentro.domain.names import display_name, normalize_name
from botcentro.domain.project_ids import normalize_year
from botcentro.errors import FailureKind, FetchError
from botcentro.http.fetcher import Fetched, NotModified, SafeFetcher

BASE = "https://www.camara.gov.co"
PAGE_URL = f"{BASE}/proyectos-de-ley/"
AJAX_URL = f"{BASE}/wp-admin/admin-ajax.php"
ALLOWED_DOMAINS = ("www.camara.gov.co",)
REF = "camara-pl"
XLSX = frozenset({"application/zip"})
JSON = frozenset({"application/json"})
PAGES_PER_BATCH = 20

NUMBER_RE = re.compile(r"(\d{1,5})\s*/\s*(\d{4}|\d{2})\s*([cCsS])?(?![\w])")
NONCE_RE = re.compile(r'PL_NONCE\s*:\s*"([0-9a-f]+)"')
XLSX_COLUMNS = ("No.", "No. Cámara", "No. Senado", "Fecha Cámara", "Fecha Senado", "Proyecto", "Título",
                "Objeto del proyecto", "Observaciones", "Tipo de Ley", "Autores", "Estado de Ley", "Origen de Ley",
                "Comisión(es)", "Legislatura", "Link del Proyecto")
FILTERS = {"term": "", "comision": "", "tipo": "All", "estado": "All", "origen": "All", "legislatura": "All",
           "ley_numero": "", "ley_fecha": ""}


def parse_numbers(raw: Any) -> list[tuple[str, int | None, str | None]]:
    """«396/2026C ACUM 120/2025» → [("396", 2026, "C"), ("120", 2025, None)]. Enteros sin año: año None."""
    if raw is None:
        return []
    if isinstance(raw, (int, float)):
        return [(str(int(raw)), None, None)]
    text = str(raw)
    found = [(str(int(n)), normalize_year(y), (s or "").upper() or None) for n, y, s in NUMBER_RE.findall(text)]
    if not found and text.strip().isdigit():
        found = [(str(int(text.strip())), None, None)]
    return found


def format_ref(initiative: str, number: str, year: int | None, corporation: str) -> str | None:
    """Misma forma que domain.project_ids.format_ref; None si falta el año (identidad incompleta)."""
    if year is None:
        return None
    prefix = "PAL" if initiative == "proyecto_acto_legislativo" else "PL"
    return f"{prefix} {number}/{year} {'Senado' if corporation == 'senado' else 'Cámara'}"


def project_key(raw_camara: Any, link: str | None) -> str | None:
    numbers = parse_numbers(raw_camara)
    if numbers and numbers[0][1] is not None:
        number, year, _ = numbers[0]
        return f"{REF}:project:{number}-{year}"
    if link:
        return f"{REF}:project:slug:{link.rstrip('/').rsplit('/', 1)[-1]}"
    return None


def _date(value: Any) -> PartialDate:
    if isinstance(value, datetime):
        return PartialDate.of_day(value.date())
    if isinstance(value, date):
        return PartialDate.of_day(value)
    if isinstance(value, str) and re.fullmatch(r"\d{4}-\d{2}-\d{2}", value.strip()):
        return PartialDate.of_day(date.fromisoformat(value.strip()))
    return PartialDate.unknown()


def _text(value: Any) -> str | None:
    if value is None:
        return None
    text = " ".join(str(value).split())
    return text or None


class CamaraProyectosConnector:
    adapter = "camara_proyectos"
    version = "0.1.0"
    parser_version = "camara-pl-parser-1"

    def __init__(self, fetcher: SafeFetcher, *, per_page: int = 50) -> None:
        self._fetcher = fetcher
        self._per_page = per_page
        self._nonce: str | None = None
        self._total_pages: int | None = None

    def validate_source(self, config: Mapping[str, Any]) -> ValidationReport:
        return ValidationReport(capabilities=["project", "status", "authorship"], diagnostics=[])

    # -- nonce público de la página ------------------------------------------------------------
    def _refresh_nonce(self) -> str:
        page = self._fetcher.fetch(PAGE_URL, accept_mimes=frozenset({"text/html"}))
        assert isinstance(page, Fetched)
        match = NONCE_RE.search(page.content.decode("utf-8", "replace"))
        if not match:
            raise FetchError("NONCE_NOT_FOUND", "la página de proyectos no publica el nonce esperado",
                             kind=FailureKind.SCHEMA_CHANGED)
        self._nonce = match.group(1)
        return self._nonce

    def _post(self, data: dict[str, Any], accept: frozenset[str]) -> Fetched:
        for attempt in (1, 2):
            nonce = self._nonce or self._refresh_nonce()
            try:
                result = self._fetcher.fetch(AJAX_URL, method="POST", accept_mimes=accept,
                                             data={**{k: str(v) for k, v in data.items()}, "_ajax_nonce": nonce})
            except FetchError as exc:
                if attempt == 1 and exc.status in (400, 403):  # nonce vencido: se renueva una vez
                    self._nonce = None
                    continue
                raise
            assert isinstance(result, Fetched)
            return result
        raise AssertionError("inalcanzable")

    def _page_count(self) -> int:
        if self._total_pages is None:
            first = self._post({"action": "get_proyectos_ley_page", "page": 1, "per_page": self._per_page,
                                **FILTERS, "comision_adv": "All"}, JSON)
            self._total_pages = int(json.loads(first.content)["data"]["total_pages"])
        return self._total_pages

    # -- §7.1 ----------------------------------------------------------------------------------
    def discover(self, cursor: Cursor) -> DiscoverPage:
        batch = int(cursor.position.get("batch", 0))
        if batch == 0:
            item = DiscoveredItem("listado_xlsx", "listado-xlsx", AJAX_URL, accept_mimes=XLSX,
                                  hints={"action": "download_proyectos_ley_xlsx"})
            return DiscoverPage([item], cursor.advance(batch=1), has_more=True)
        total = self._page_count()
        first = (batch - 1) * PAGES_PER_BATCH + 1
        last = min(first + PAGES_PER_BATCH - 1, total)
        items = [DiscoveredItem("ajax_page", f"ajax-page:{n}", AJAX_URL, accept_mimes=JSON, hints={"page": n})
                 for n in range(first, last + 1)]
        return DiscoverPage(items, cursor.advance(batch=batch + 1), has_more=last < total)

    def fetch(self, item: DiscoveredItem, *, etag: str | None = None,
              last_modified: str | None = None) -> Fetched | NotModified:
        if item.record_type == "listado_xlsx":
            return self._post({"action": "download_proyectos_ley_xlsx", **FILTERS}, XLSX)
        return self._post({"action": "get_proyectos_ley_page", "page": item.hints["page"],
                           "per_page": self._per_page, **FILTERS, "comision_adv": "All"}, JSON)

    def parse(self, item: DiscoveredItem, snapshot: Fetched) -> ParseResult:
        if item.record_type == "listado_xlsx":
            return self._parse_xlsx(snapshot.content)
        return self._parse_page(snapshot.content)

    # -- XLSX ----------------------------------------------------------------------------------
    def _parse_xlsx(self, content: bytes) -> ParseResult:
        import openpyxl

        try:
            sheet = openpyxl.load_workbook(io.BytesIO(content)).active
        except Exception as exc:  # noqa: BLE001 — cualquier fallo del formato es cambio de esquema
            return ParseResult(issues=[Issue("INVALID_XLSX", str(exc)[:200], IssueSeverity.SCHEMA_CHANGE)])
        header = tuple(c.value for c in sheet[1])
        if header != XLSX_COLUMNS:
            return ParseResult(issues=[Issue("COLUMNS_CHANGED", f"columnas: {header}", IssueSeverity.SCHEMA_CHANGE)])

        candidates: list[NormalizedCandidate] = []
        issues: list[Issue] = []
        for row_index in range(2, sheet.max_row + 1):
            cells = [sheet.cell(row_index, c) for c in range(1, len(XLSX_COLUMNS) + 1)]
            row = dict(zip(XLSX_COLUMNS, (c.value for c in cells)))
            link = cells[15].hyperlink.target if cells[15].hyperlink else None
            pointer = f"/fila/{row_index}"
            key = project_key(row["No. Cámara"], link)
            if key is None:
                issues.append(Issue("NO_IDENTITY", "fila sin número de Cámara ni enlace", IssueSeverity.QUARANTINE, pointer))
                continue
            candidates.extend(self._row_candidates(row, link, key, pointer))
        return ParseResult(candidates=candidates, issues=issues)

    def _row_candidates(self, row: Mapping[str, Any], link: str | None, key: str, pointer: str) -> list[NormalizedCandidate]:
        law_type = _text(row["Tipo de Ley"])
        initiative = "proyecto_acto_legislativo" if law_type == "Acto Legislativo" else "proyecto_ley"
        camara_numbers = parse_numbers(row["No. Cámara"])
        senado_numbers = parse_numbers(row["No. Senado"])

        camara_ref = None
        accumulated: list[str] = []
        if camara_numbers:
            number, year, _ = camara_numbers[0]
            camara_ref = format_ref(initiative, number, year, "camara")
            accumulated = [r for n, y, _ in camara_numbers[1:] if (r := format_ref(initiative, n, y, "camara"))]

        senado_ref = None
        quarantine = None
        if senado_numbers:
            number, year, suffix = senado_numbers[0]
            if suffix == "C":
                quarantine = f"número en la columna Senado con sufijo de Cámara: {row['No. Senado']!r}"
            else:
                senado_ref = format_ref(initiative, number, year, "senado")

        filed_camara, filed_senado = _date(row["Fecha Cámara"]), _date(row["Fecha Senado"])
        profile = NormalizedCandidate(
            subject_type="project", subject_ref=key, predicate="project_profile",
            value={
                "title": _text(row["Título"]), "short_name": _text(row["Proyecto"]),
                "object": _text(row["Objeto del proyecto"]), "observations": _text(row["Observaciones"]),
                "law_type": law_type, "initiative_type": initiative, "origin": _text(row["Origen de Ley"]),
                "commissions": _text(row["Comisión(es)"]), "legislature": _text(row["Legislatura"]),
                "link": link, "camara_raw": _text(row["No. Cámara"]), "senado_raw": _text(row["No. Senado"]),
                "camara_ref": camara_ref, "senado_ref": senado_ref, "accumulated_refs": accumulated,
                "explicit_link": bool(camara_ref and senado_ref),
                "filed_camara": filed_camara.iso(), "filed_senado": filed_senado.iso(),
                "authors_raw": _text(row["Autores"]),
            },
            value_raw=_text(row["Título"]), effective=filed_camara, record_pointer=pointer,
            quarantine_reason=quarantine,
        )
        candidates = [profile]
        status = _text(row["Estado de Ley"])
        if status:
            candidates.append(NormalizedCandidate(
                subject_type="project", subject_ref=key, predicate="status_reported", value={"status_raw": status},
                value_raw=status, record_pointer=pointer,
            ))
        return candidates

    # -- Páginas JSON (autores con ID) -------------------------------------------------------------
    def _parse_page(self, content: bytes) -> ParseResult:
        try:
            body = json.loads(content)
            items: Sequence[Mapping[str, Any]] = body["data"]["items"]
            self._total_pages = int(body["data"].get("total_pages") or self._total_pages or 0)
        except (ValueError, KeyError, TypeError) as exc:
            return ParseResult(issues=[Issue("INVALID_PAGE", str(exc)[:200], IssueSeverity.SCHEMA_CHANGE)])
        if not items:
            return ParseResult(empty_source=True)

        candidates: list[NormalizedCandidate] = []
        issues: list[Issue] = []
        for index, item in enumerate(items):
            pointer = f"/data/items/{index}"
            key = project_key(item.get("nro_camara"), item.get("link_web"))
            if key is None:
                issues.append(Issue("NO_IDENTITY", "proyecto sin número ni enlace", IssueSeverity.QUARANTINE, pointer))
                continue
            for entry in filter(None, (item.get("autores_pack") or "").split("::")):
                parts = entry.split("||")
                if len(parts) != 3 or not parts[0].strip().isdigit():
                    issues.append(Issue("AUTHOR_FORMAT", f"autor no interpretable: {entry[:80]!r}",
                                        IssueSeverity.WARNING, pointer))
                    continue
                rep_id, name, slug = parts[0].strip(), display_name(parts[1]), parts[2].strip()
                person_ref = f"camara:representative:{rep_id}"
                candidates.append(NormalizedCandidate(
                    subject_type="person", subject_ref=person_ref, predicate="representative_seen",
                    value={"name": name, "normalized_name": normalize_name(name), "profile_path": slug},
                    value_raw=parts[1], record_pointer=pointer,
                ))
                candidates.append(NormalizedCandidate(
                    subject_type="authorship", subject_ref=f"{key}:author:{rep_id}", predicate="authorship",
                    value={"project_ref": key, "person_ref": person_ref, "role": "autor"},
                    value_raw=parts[1], record_pointer=pointer,
                ))
        unique = {(c.subject_ref, c.predicate, json.dumps(c.value, sort_keys=True, ensure_ascii=False)): c
                  for c in candidates}
        return ParseResult(candidates=list(unique.values()), issues=issues)
